"""
S_analyze - Auswertung der blinden Annotation (LLM gegen Kosinus, gewichtet nach Schicht).

Vorab festgelegte Hauptfrage (Praeregistrierung durch dieses Skript):
  Differenz der F1-Werte fuer die K-Entscheidung, LLM minus Kosinus bei gleicher
  Aufnahmequote (t_eq aus S_sample.py), 95-%-Bootstrap-Intervall innerhalb der Schichten.
Sekundaer: Kosinus mit auf G_manual kalibrierter Schwelle, schwellenfreie AUC (LLM-Konfidenz
gegen Kosinus), Genauigkeit der K-Entscheidung, Precision der gelieferten Mappings,
Relationsklassifikation (gewichtetes Macro-F1, Richtungsgenauigkeit), Cohen's Kappa.

Schaetzer: Horvitz-Thompson, Gewicht je Paar = N_h / n_h (Schichtgroesse durch Stichprobe).

Verwendung:
  python eval/S_analyze.py                      # nach der Annotation
  python eval/S_analyze.py --selftest           # Test mit simulierten Annotationen
"""
from __future__ import annotations

import argparse
import json
import random
import sys
from pathlib import Path

import numpy as np

ROOT = Path(__file__).resolve().parents[1]
ANN = ROOT / "eval" / "annotation"
KEY = ANN / "sample_key_NICHT_VOR_ANNOTATION_OEFFNEN.json"
OUT = ROOT / "eval" / "results" / "annotation" / "S_report.json"
K = {"EQUIVALENT", "A_SUBSUMES_B", "B_SUBSUMES_A"}
RELS = ["EQUIVALENT", "A_SUBSUMES_B", "B_SUBSUMES_A", "RELATED"]


def read_sheet(path: Path) -> tuple[dict[str, str], dict[str, str]]:
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True)["Annotation"]
    head = [c.value for c in ws[1]]
    ix = {h: i for i, h in enumerate(head)}
    gold, unsure = {}, {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[ix["ID"]]:
            g = (row[ix["Beziehung"]] or "").strip()
            gold[row[ix["ID"]]] = g
            unsure[row[ix["ID"]]] = (row[ix["Unsicher"]] or "").strip()
    return gold, unsure


def prf(d, y, w):
    tp, fp = float(np.sum(w * d * y)), float(np.sum(w * d * (1 - y)))
    fn, tn = float(np.sum(w * (1 - d) * y)), float(np.sum(w * (1 - d) * (1 - y)))
    p = tp / (tp + fp) if tp + fp else 0.0
    r = tp / (tp + fn) if tp + fn else 0.0
    f = 2 * p * r / (p + r) if p + r else 0.0
    return p, r, f, (tp + tn) / (tp + fp + fn + tn)


def wauc(s, y, w):
    pos, neg = y == 1, y == 0
    if not pos.any() or not neg.any():
        return float("nan")
    a, b = s[pos][:, None], s[neg][None, :]
    ww = w[pos][:, None] * w[neg][None, :]
    return float((((a > b) + 0.5 * (a == b)) * ww).sum() / ww.sum())


def rel_metrics(pred, gold, w):
    W = {g: {} for g in RELS}
    for p, g, x in zip(pred, gold, w):
        W[g][p] = W[g].get(p, 0.0) + float(x)
    f1 = {}
    for c in RELS:
        tp = W[c].get(c, 0.0)
        fn = sum(W[c].values()) - tp
        fp = sum(W[g].get(c, 0.0) for g in RELS if g != c)
        f1[c] = 2 * tp / (2 * tp + fp + fn) if 2 * tp + fp + fn else 0.0
    tot = sum(sum(v.values()) for v in W.values())
    acc = sum(W[c].get(c, 0.0) for c in RELS) / tot
    sub = [(p, g, x) for p, g, x in zip(pred, gold, w) if g in ("A_SUBSUMES_B", "B_SUBSUMES_A") and p in ("A_SUBSUMES_B", "B_SUBSUMES_A")]
    dacc = sum(x for p, g, x in sub if p == g) / sum(x for _, _, x in sub) if sub else float("nan")
    return {"macro_f1": float(np.mean(list(f1.values()))), "per_class_f1": f1, "accuracy": acc, "direction_accuracy": dacc, "n_direction": len(sub)}


def kappa(a, b):
    n = len(a)
    po = sum(x == y for x, y in zip(a, b)) / n
    pe = sum((a.count(l) / n) * (b.count(l) / n) for l in set(a) | set(b))
    return (po - pe) / (1 - pe) if pe < 1 else float("nan")


def analyse(key: dict, gold: dict[str, str], second: dict[str, str] | None, n_boot: int = 4000, seed: int = 20261005) -> dict:
    ids = sorted(key["pairs"])
    P = [key["pairs"][i] for i in ids]
    g = [gold[i] for i in ids]
    y = np.array([x in K for x in g], dtype=float)
    d = np.array([p["d"] for p in P], dtype=float)
    c = np.array([p["c"] for p in P], dtype=float)
    w = np.array([p["weight"] for p in P], dtype=float)
    sim = np.array([p["sim"] for p in P], dtype=float)
    conf = np.array([p["conf"] for p in P], dtype=float)
    llm_score = np.where(d == 1, conf, -conf)
    tg = key["meta"].get("t_gold")
    cg = (sim >= tg).astype(float) if tg is not None else None
    strata = np.array([p["stratum"] for p in P])
    groups = [np.where(strata == s)[0] for s in sorted(set(strata))]
    rng = np.random.default_rng(seed)

    def stats(idx):
        yy, ww = y[idx], w[idx]
        out = {"f1_llm": prf(d[idx], yy, ww)[2], "f1_cos": prf(c[idx], yy, ww)[2], "f1_all": prf(np.ones(len(idx)), yy, ww)[2],
               "acc_llm": prf(d[idx], yy, ww)[3], "acc_cos": prf(c[idx], yy, ww)[3],
               "prec_llm": prf(d[idx], yy, ww)[0], "auc_llm": wauc(llm_score[idx], yy, ww), "auc_cos": wauc(sim[idx], yy, ww)}
        if cg is not None:
            out["f1_cos_gold"] = prf(cg[idx], yy, ww)[2]
        return out

    point = stats(np.arange(len(ids)))
    boots = []
    for _ in range(n_boot):
        idx = np.concatenate([rng.choice(gr, size=len(gr), replace=True) for gr in groups])
        boots.append(stats(idx))

    def ci(name_a, name_b=None):
        v = np.array([b[name_a] - (b[name_b] if name_b else 0.0) for b in boots])
        v = v[~np.isnan(v)]
        est = point[name_a] - (point[name_b] if name_b else 0.0)
        p2 = float(2 * min((v <= 0).mean(), (v >= 0).mean())) if name_b else None
        return {"estimate": float(est), "ci95": [float(np.percentile(v, 2.5)), float(np.percentile(v, 97.5))],
                **({"p_bootstrap_two_sided": min(1.0, p2)} if name_b else {})}

    p_llm, r_llm, f_llm, a_llm = prf(d, y, w)
    p_cos, r_cos, f_cos, a_cos = prf(c, y, w)
    res = {
        "n_annotated": len(ids),
        "gold_distribution_unweighted": {r: g.count(r) for r in RELS},
        "gold_k_share_estimated": float(np.sum(w * y) / np.sum(w)),
        "primary_delta_f1_llm_minus_cosine_equal_rate": ci("f1_llm", "f1_cos"),
        "llm": {"precision": p_llm, "recall": r_llm, "f1": f_llm, "accuracy": a_llm},
        "cosine_equal_rate": {"precision": p_cos, "recall": r_cos, "f1": f_cos, "accuracy": a_cos, "threshold": key["meta"]["t_eq"]},
        "all_in_k_f1": point["f1_all"],
        "precision_of_delivered_mappings": ci("prec_llm"),
        "secondary_delta_accuracy": ci("acc_llm", "acc_cos"),
        "secondary_auc_llm_confidence": ci("auc_llm"),
        "secondary_auc_cosine": ci("auc_cos"),
        "secondary_delta_auc": ci("auc_llm", "auc_cos"),
        "relation_classification_llm": rel_metrics([p["rel"] for p in P], g, w),
        "n_boot": n_boot, "seed": seed,
    }
    if cg is not None:
        res["secondary_cosine_gold_threshold"] = {"threshold": tg, "f1": point["f1_cos_gold"]}
        res["secondary_delta_f1_vs_gold_threshold"] = ci("f1_llm", "f1_cos_gold")
    lo, hi = res["primary_delta_f1_llm_minus_cosine_equal_rate"]["ci95"]
    res["verdict_primary"] = ("LLM gesichert besser" if lo > 0 else "Kosinus gesichert besser" if hi < 0 else "kein gesicherter Unterschied")
    if second:
        common = [i for i in ids if second.get(i)]
        a, b = [gold[i] for i in common], [second[i] for i in common]
        res["interrater"] = {"n": len(common), "agreement": sum(x == y_ for x, y_ in zip(a, b)) / len(common),
                             "kappa_4_classes": kappa(a, b),
                             "kappa_k_membership": kappa([x in K for x in a], [x in K for x in b])}
    return res


def selftest(key: dict) -> None:
    rnd = random.Random(1)
    pk = {"KK": 0.9, "RR": 0.05, "LLM_nur": 0.55, "Kosinus_nur": 0.2}
    gold, second = {}, {}
    for i, p in key["pairs"].items():
        if rnd.random() < pk[p["stratum"]]:
            gold[i] = p["rel"] if p["rel"] in K and rnd.random() < 0.5 else rnd.choice(sorted(K))
        else:
            gold[i] = "RELATED"
        if p["second"]:
            second[i] = gold[i] if rnd.random() < 0.85 else rnd.choice(RELS)
    res = analyse(key, gold, second, n_boot=1000)
    print("SELBSTTEST (simulierte Annotation, nur zum Pruefen des Codes, nicht inhaltlich lesen):")
    for k in ("n_annotated", "primary_delta_f1_llm_minus_cosine_equal_rate", "verdict_primary", "llm", "cosine_equal_rate",
              "secondary_delta_auc", "interrater"):
        print(" ", k, "->", json.dumps(res.get(k), ensure_ascii=False, default=float)[:260])


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--selftest", action="store_true")
    ap.add_argument("--sheet", default=str(ANN / "annotation_sheet.xlsx"))
    ap.add_argument("--second", default=str(ANN / "annotation_second.xlsx"))
    args = ap.parse_args()
    key = json.loads(KEY.read_text(encoding="utf-8"))
    if args.selftest:
        selftest(key)
        return
    gold, unsure = read_sheet(Path(args.sheet))
    bad = [i for i in key["pairs"] if gold.get(i) not in RELS]
    if bad:
        raise SystemExit(f"{len(bad)} Paare ohne gueltige Beziehung, z. B. {bad[:8]}")
    second = None
    if Path(args.second).exists():
        s2, _ = read_sheet(Path(args.second))
        second = {i: v for i, v in s2.items() if v in RELS} or None
    res = analyse(key, gold, second)
    res["unsure_share"] = sum(v.lower() == "ja" for v in unsure.values()) / len(unsure)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(res, indent=2, ensure_ascii=False, default=float), encoding="utf-8")
    print(json.dumps(res, indent=2, ensure_ascii=False, default=float))
    print("Bericht:", OUT)


if __name__ == "__main__":
    main()
