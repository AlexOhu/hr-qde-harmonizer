"""
S_sample - geschichtete Stichprobe aus den 800 Lieferpaaren (v0.4) fuer die blinde Annotation.

Ziel: LLM-Entscheidung "Paar gehoert zu K" gegen eine Kosinus-Baseline statistisch absichern.
Die Stichprobe ist nach der Uebereinstimmung beider Entscheidungen geschichtet, weil nur
Paare mit unterschiedlicher Entscheidung den Unterschied tragen:

  KK          LLM nimmt auf, Kosinus nimmt auf
  RR          LLM verwirft,  Kosinus verwirft
  LLM_nur     LLM nimmt auf, Kosinus verwirft
  Kosinus_nur LLM verwirft,  Kosinus nimmt auf

Kosinus-Schwelle (vorab festgelegt, ohne Gold): Gleiche Aufnahmequote wie das LLM
(t_eq = die n-te hoechste Aehnlichkeit, n = Zahl der vom LLM aufgenommenen Paare).
Zusaetzlich wird eine auf G_manual kalibrierte Schwelle protokolliert (nur Sekundaeranalyse).

Erzeugt in eval/annotation/:
  annotation_sheet.xlsx                         Liste zum Annotieren, OHNE Modellausgabe
  annotation_second.xlsx                        Teilmenge fuer eine zweite Person (optional)
  sample_key_NICHT_VOR_ANNOTATION_OEFFNEN.json  Schluessel mit Modellausgabe, Schicht, Gewicht

Verwendung:  python eval/S_sample.py [--n-kk 20 --n-rr 20 --n-llm-only 30 --n-cos-only 30 --n-second 30]
"""
from __future__ import annotations

import argparse
import glob
import json
import random
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

K = {"EQUIVALENT", "A_SUBSUMES_B", "B_SUBSUMES_A"}
OUT = ROOT / "eval" / "annotation"
RELATIONS = ["EQUIVALENT", "A_SUBSUMES_B", "B_SUBSUMES_A", "RELATED"]

GUIDE = [
    ("Aufgabe", "Bestimme fuer jedes Paar die Beziehung zwischen dem e-CF-Konzept (A) und dem ESCO-Skill (B). Bitte in der vorgegebenen Reihenfolge arbeiten und keine Modellausgaben, Mapping-Dateien oder Ergebnisse der Pipeline ansehen."),
    ("Richtung", "A ist immer das e-CF-Konzept, B immer der ESCO-Skill. \"X ist Oberbegriff von Y\" bedeutet: Y ist ein Spezialfall von X, jede Instanz von Y ist auch eine Instanz von X. Y ist das engere, X das weitere Konzept."),
    ("EQUIVALENT", "Konzept A und Konzept B bezeichnen dasselbe Konzept."),
    ("A_SUBSUMES_B", "Konzept A ist Oberbegriff von Konzept B. Das heisst: B ist enger/spezifischer, A ist weiter/allgemeiner."),
    ("B_SUBSUMES_A", "Konzept B ist Oberbegriff von Konzept A. Das heisst: A ist enger/spezifischer, B ist weiter/allgemeiner."),
    ("RELATED", "Thematisch verwandt, aber weder aequivalent, noch ist eines der beiden Konzepte Oberbegriff des anderen."),
    ("Vorgehen", "Zuerst pruefen, welches der beiden Konzepte das engere ist, und daraus die Beziehung ableiten. Im Zweifel zwischen RELATED und Subsumtion bitte die Spalte Unsicher mit ja markieren."),
    ("ESCO-Skills", "Das ESCO-Skelett enthaelt keine Definitionen. Im Zweifel den Skill ueber den Link in der Spalte ESCO-Link im ESCO-Portal nachschlagen. Pipeline-Ergebnisse und Mapping-Dateien bleiben tabu."),
    ("Zeitbedarf", "Etwa 2 bis 3 Minuten je Paar."),
]


def load_results(project_glob: str) -> list[dict]:
    d = sorted(glob.glob(str(ROOT / "data" / "projects" / project_glob)))
    if not d:
        raise SystemExit(f"Kein Projekt gefunden: {project_glob}")
    r = json.loads((Path(d[-1]) / "classification_results.json").read_text(encoding="utf-8"))
    return r if isinstance(r, list) else r.get("results", r)


def gold_threshold() -> float | None:
    """Kosinus-Schwelle mit bestem F1 fuer K auf G_manual (Sekundaeranalyse)."""
    try:
        from eval.B_baselines import b1_cosine_k, cosine_scores
        f = sorted((ROOT / "eval/results/rerun_2026-10").glob("prod0_nfs0_t0.0_pass1.json"))[0]
        rows = json.loads(f.read_text(encoding="utf-8"))["pairs"]
        scores = cosine_scores(rows)
        return float(b1_cosine_k(scores, [r["gold"] for r in rows])["B1_cosine_k_best"]["threshold"])
    except Exception as exc:  # Modell oder Daten nicht verfuegbar
        print("Hinweis: G_manual-Schwelle nicht berechnet:", exc)
        return None


def write_sheet(path: Path, rows: list[dict], title: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Annotation"
    head = ["ID", "e-CF-Konzept (A)", "e-CF-Definition", "ESCO-Skill (B)", "ESCO-Link",
            "Beziehung", "Unsicher", "Kommentar"]
    ws.append(head)
    for r in rows:
        ws.append([r["id"], r["a_label"], r["a_def"], r["b_label"], r["b"], None, None, None])
    widths = [7, 34, 60, 34, 36, 18, 10, 40]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE6F0")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for col in ("F", "G"):
        for r in range(2, len(rows) + 2):
            ws[f"{col}{r}"].fill = PatternFill("solid", fgColor="FFF9DB")
    ws.freeze_panes = "B2"
    dv = DataValidation(type="list", formula1='"' + ",".join(RELATIONS) + '"', allow_blank=True)
    dv.error, dv.errorTitle = "Bitte einen Wert aus der Liste waehlen.", "Ungueltiger Wert"
    ws.add_data_validation(dv)
    dv.add(f"F2:F{len(rows) + 1}")
    dv2 = DataValidation(type="list", formula1='"ja,nein"', allow_blank=True)
    ws.add_data_validation(dv2)
    dv2.add(f"G2:G{len(rows) + 1}")
    g = wb.create_sheet("Anleitung")
    g.append([title])
    g["A1"].font = Font(bold=True)
    g.append([])
    for k, v in GUIDE:
        g.append([k, v])
    g.column_dimensions["A"].width = 16
    g.column_dimensions["B"].width = 120
    for row in g.iter_rows(min_row=3):
        row[0].font = Font(bold=True)
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--project", default="*v0_4*")
    ap.add_argument("--n-kk", type=int, default=20)
    ap.add_argument("--n-rr", type=int, default=20)
    ap.add_argument("--n-llm-only", type=int, default=30)
    ap.add_argument("--n-cos-only", type=int, default=30)
    ap.add_argument("--n-second", type=int, default=30)
    ap.add_argument("--seed", type=int, default=20261005)
    args = ap.parse_args()

    from eval.E2_few_shot import _load_concept_lookup
    lookup = _load_concept_lookup()
    results = load_results(args.project)
    pairs = []
    for r in results:
        p = r["pair"]
        rel = r.get("final_relation") or "NO_CONSENSUS"
        pairs.append({"a": p["concept_a"], "b": p["concept_b"], "sim": float(p["similarity_score"]),
                      "rel": rel, "conf": float(r.get("mapping_confidence") or 0.0), "d": rel in K})
    pairs.sort(key=lambda x: (x["a"], x["b"]))
    n_acc = sum(p["d"] for p in pairs)
    t_eq = sorted((p["sim"] for p in pairs), reverse=True)[n_acc - 1]
    for p in pairs:
        p["c"] = p["sim"] >= t_eq
        p["stratum"] = ("KK" if p["d"] and p["c"] else "RR" if not p["d"] and not p["c"]
                        else "LLM_nur" if p["d"] else "Kosinus_nur")
    N = Counter(p["stratum"] for p in pairs)
    want = {"KK": args.n_kk, "RR": args.n_rr, "LLM_nur": args.n_llm_only, "Kosinus_nur": args.n_cos_only}
    rng = random.Random(args.seed)
    sample = []
    for s, n in want.items():
        pool = [p for p in pairs if p["stratum"] == s]
        take = rng.sample(pool, min(n, len(pool)))
        for p in take:
            p["N_h"], p["n_h"] = len(pool), len(take)
            p["weight"] = len(pool) / len(take)
        sample += take
    rng.shuffle(sample)
    for i, p in enumerate(sample, 1):
        p["id"] = f"P{i:03d}"
        la, lb = lookup.get(p["a"], {}), lookup.get(p["b"], {})
        p.update(a_label=la.get("pref_label", p["a"]), a_def=la.get("definition") or "",
                 b_label=lb.get("pref_label", p["b"]), b_def=lb.get("definition") or "")
    second_ids = {p["id"] for p in random.Random(args.seed + 1).sample(sample, min(args.n_second, len(sample)))}
    for p in sample:
        p["second"] = p["id"] in second_ids

    OUT.mkdir(parents=True, exist_ok=True)
    write_sheet(OUT / "annotation_sheet.xlsx", sample, "Blinde Annotation: e-CF gegen ESCO (Stichprobe aus der Lieferung v0.4)")
    sec = [p for p in sample if p["second"]]
    random.Random(args.seed + 2).shuffle(sec)
    write_sheet(OUT / "annotation_second.xlsx", sec, "Zweitannotation (Teilmenge), blind und unabhaengig vom Erstannotator")
    key = {"meta": {"seed": args.seed, "n_pairs_total": len(pairs), "n_accepted_llm": n_acc, "t_eq": t_eq,
                    "t_gold": gold_threshold(), "strata_N": dict(N), "strata_n": {s: sum(p["stratum"] == s for p in sample) for s in want},
                    "project": args.project},
           "pairs": {p["id"]: {k: p[k] for k in ("a", "b", "a_label", "b_label", "sim", "rel", "conf", "d", "c",
                                                  "stratum", "N_h", "n_h", "weight", "second")} for p in sample}}
    (OUT / "sample_key_NICHT_VOR_ANNOTATION_OEFFNEN.json").write_text(json.dumps(key, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"Paare gesamt {len(pairs)}, vom LLM aufgenommen {n_acc}, Kosinus-Schwelle (gleiche Quote) t_eq = {t_eq:.4f}")
    print("Schichten (N gesamt / gezogen):", {s: f"{N[s]} / {key['meta']['strata_n'][s]}" for s in want})
    print(f"Stichprobe: {len(sample)} Paare, davon {len(sec)} fuer die Zweitannotation")
    print("G_manual-Schwelle (Sekundaeranalyse):", key["meta"]["t_gold"])
    print("Ausgabe:", OUT)


if __name__ == "__main__":
    main()
